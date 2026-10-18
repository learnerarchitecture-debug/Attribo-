"""
Document parsing: docx, pdf, txt, md, rtf, odt.

(c) 2026 Attribo. All rights reserved.
"""
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = {".docx", ".pdf", ".txt", ".md", ".rtf", ".odt"}


@dataclass
class ParsedDoc:
    filename: str = ""
    path: str = ""
    full_text: str = ""
    pages: list = field(default_factory=list)
    warnings: list = field(default_factory=list)


def _read_text_file(path) -> str:
    raw = Path(path).read_bytes()
    for enc in ("utf-8", "utf-16", "cp1256", "latin-1"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return raw.decode("utf-8", errors="ignore")


def _parse_docx(path) -> str:
    import docx

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _parse_pdf(path) -> list:
    from pypdf import PdfReader

    reader = PdfReader(path)
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return pages


def _strip_rtf(text: str) -> str:
    text = re.sub(
        r"\\'([0-9a-fA-F]{2})",
        lambda m: bytes([int(m.group(1), 16)]).decode("latin-1", "ignore"),
        text,
    )
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
    text = text.replace("{", " ").replace("}", " ")
    return re.sub(r"[ \t]+", " ", text)


def _parse_odt(path) -> str:
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("content.xml").decode("utf-8", errors="ignore")
    xml = xml.replace("</text:p>", "\n")
    return re.sub(r"<[^>]+>", " ", xml)


def parse_document(path) -> ParsedDoc:
    p = Path(path)
    doc = ParsedDoc(filename=p.name, path=str(p))
    ext = p.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        doc.warnings.append(f"Unsupported extension '{ext}' - attempting plain text read.")
        doc.full_text = _read_text_file(path)
        return doc
    try:
        if ext == ".docx":
            text = _parse_docx(path)
            doc.pages = [text]
        elif ext == ".pdf":
            pages = _parse_pdf(path)
            doc.pages = pages
            text = "\n".join(pages)
            if len(text.strip()) < 50:
                doc.warnings.append(
                    "This PDF appears to be scanned (no text layer). OCR is not "
                    "included in this build - results will be limited."
                )
        elif ext in (".txt", ".md"):
            text = _read_text_file(path)
            doc.pages = [text]
        elif ext == ".rtf":
            text = _strip_rtf(_read_text_file(path))
            doc.pages = [text]
        elif ext == ".odt":
            text = _parse_odt(path)
            doc.pages = [text]
        else:  # pragma: no cover
            text = _read_text_file(path)
            doc.pages = [text]
    except Exception as exc:
        raise RuntimeError(f"Could not parse '{p.name}': {exc}") from exc
    doc.full_text = text or ""
    if not doc.full_text.strip():
        doc.warnings.append("No extractable text found in the document.")
    return doc
