"""
Offline smoke test for Attribo.

Run from the project root:
    C:\\apcvenv\\Scripts\\python.exe tests\\smoke_test.py

No network access required.

(c) 2026 Attribo. All rights reserved.
"""
import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def main() -> int:
    from attribo.storage import db

    db.init_db()
    print("[ok] database init")

    # -- matching ------------------------------------------------------
    from attribo.core.matching import (
        content_stems,
        find_near_runs,
        find_verbatim_runs,
        paraphrase_check,
        tokenize_spans,
    )

    a = "The quick brown fox jumps over the lazy dog near the river bank in the morning light"
    b = "The quick brown fox jumps over the lazy dog and then it went home"
    runs = find_verbatim_runs(tokenize_spans(a), [t[0] for t in tokenize_spans(b)], 8)
    assert runs and runs[0][3] >= 8, "verbatim detection failed"
    print("[ok] verbatim matching")

    a2 = (
        "urban planning policies strongly influence residential development "
        "patterns across large metropolitan regions"
    )
    b2 = (
        "urban planning strategies strongly influence residential development "
        "patterns across large metropolitan regions"
    )
    near = find_near_runs(tokenize_spans(a2), [t[0] for t in tokenize_spans(b2)], 8)
    assert near and near[0][3] >= 8, "near-match detection failed"
    print("[ok] near-match matching")

    from attribo.core.matching import (
        TfidfIndex,
        content_stem_list,
        find_internal_duplicates,
    )

    part = "alpha bravo charlie delta echo foxtrot golf hotel india juliet kilo lima"
    filler = " ".join(f"filler{i:02d}" for i in range(35))
    dups = find_internal_duplicates(tokenize_spans(f"{part} {filler} {part}"), 8, 30)
    assert dups and dups[0][3] >= 8, "internal duplicate detection failed"
    print("[ok] internal duplicate detection")

    src_a = content_stem_list(
        [t[0] for t in tokenize_spans(
            "finite element modelling software was used to carry out structural analysis of tall buildings"
        )]
    )
    src_b = content_stem_list(
        [t[0] for t in tokenize_spans(
            "soil mechanics foundations retaining walls lateral earth pressure design methods"
        )]
    )
    tfidf = TfidfIndex([src_a, src_b])
    query = content_stem_list(
        [t[0] for t in tokenize_spans(
            "structural analysis was performed using finite element modelling techniques for buildings"
        )]
    )
    hit = tfidf.best_cosine(query, min_cosine=0.2, min_shared=3)
    assert hit and hit[1] == 0, "tf-idf index picked the wrong source"
    print("[ok] TF-IDF paraphrase index")

    # -- web engine provider selection (offline) ---------------------------
    from attribo.engines.web_engine import (
        SearchHit,
        WebChecker,
        dedupe_hits,
    )

    checker = WebChecker({"brave_api_key": "x", "serper_api_key": "y"})
    assert checker.available(), "provider detection failed"
    assert checker.configured_providers() == ["brave", "serper"], "provider order wrong"
    assert WebChecker({}).search("phrase") == [], "empty checker returned hits"
    merged = dedupe_hits(
        [
            SearchHit("a", "https://example.com/x/", ""),
            SearchHit("b", "https://example.com/x", ""),
            SearchHit("c", "https://other.org/", ""),
        ]
    )
    assert [h.title for h in merged] == ["a", "c"], "URL dedupe failed"
    print("[ok] web engine provider selection")

    sentence = "structural analysis was performed using finite element modelling software"
    source = content_stems(
        [
            t[0]
            for t in tokenize_spans(
                "finite element modelling software was used to carry out structural analysis"
            )
        ]
    )
    assert paraphrase_check([t[0] for t in tokenize_spans(sentence)], source) is not None, (
        "paraphrase heuristic failed"
    )
    print("[ok] paraphrase heuristic (stemmed)")

    # -- AI heuristic --------------------------------------------------
    from attribo.core.ai_detect import analyze_ai

    ai = analyze_ai("The system was designed carefully. " * 40)
    assert "score" in ai and ai.get("version") == "stylometry-1.5"
    print("[ok] ai heuristic (stylometry-1.5)")

    # -- citation gaps / references -------------------------------------
    from attribo.core.citation import find_citation_gaps

    gaps = find_citation_gaps(
        [
            "Concrete production contributes approximately eight percent of global carbon "
            "emissions. This figure increased sharply between 2010 and 2020 across many "
            "developing regions."
        ]
    )
    assert gaps, "citation gap detection failed"
    no_gaps = find_citation_gaps(
        [
            "Concrete production contributes 8% of emissions (Smith, 2019). "
            "This figure increased between 2010 and 2020 (Jones, 2021)."
        ]
    )
    assert not no_gaps, "cited sentences were wrongly flagged"
    print("[ok] citation gap finder")

    from attribo.core.references import _guess_title, split_entries

    refs = split_entries(
        [
            "[1] Smith, J. (2019). Structural analysis of tall buildings. "
            "Journal of Structures, 12(3), 45-67.\n"
            "[2] Jones, A. and Lee, B. (2021). Concrete creep in high-rise "
            "construction. Cement and Concrete Research, 8(1), 1-20."
        ]
    )
    assert len(refs) == 2, f"reference split failed: {len(refs)}"
    assert "Structural analysis" in _guess_title(refs[0]), "title guess failed"
    print("[ok] reference parsing")

    # -- pipeline (local mode) ------------------------------------------
    tmp = Path(tempfile.mkdtemp(prefix="attribo_"))
    doc = tmp / "sample.txt"
    doc.write_text(
        "This is a sample thesis paragraph about structural engineering and "
        "construction methods.\n\n"
        + "Urban planning policies influence residential development patterns "
        "across metropolitan regions worldwide. " * 4,
        encoding="utf-8",
    )
    settings = {
        "enable_web": False,
        "enable_academic": False,
        "enable_ai": True,
        "max_queries": 5,
        "min_match_words": 8,
        "exclude_quotes": False,
        "google_api_key": "",
        "google_cx": "",
        "brave_api_key": "",
        "license_key": "",
        "fetch_timeout": 5,
        "max_fetch_per_query": 2,
    }
    from attribo.core.pipeline import run_scan

    result = run_scan(str(doc), settings)
    assert result.report_id and result.doc_hash and result.words > 0
    print("[ok] pipeline (local mode)")

    # -- local archive ----------------------------------------------------
    archive_dir = tmp / "archive"
    archive_dir.mkdir()
    (archive_dir / "previous_thesis.txt").write_text(
        "Urban planning policies influence residential development patterns "
        "across metropolitan regions worldwide. " * 4,
        encoding="utf-8",
    )
    archive_settings = dict(settings)
    archive_settings.update({"enable_archive": True, "archive_folder": str(archive_dir)})
    archive_result = run_scan(str(doc), archive_settings)
    assert any(s.kind == "archive" for s in archive_result.sources), "archive source missing"
    assert any(m.match_type in ("verbatim", "near") for m in archive_result.matches), (
        "archive match missing"
    )
    assert archive_result.similarity > 0, "archive similarity should be > 0"
    print("[ok] local archive comparison")

    # -- findings HTML -----------------------------------------------------
    from attribo.report.findings import build_findings_html

    findings_file = tmp / "findings.html"
    build_findings_html(archive_result, str(findings_file))
    assert "Attribo" in findings_file.read_text(encoding="utf-8")
    print("[ok] findings HTML export")

    # -- highlighted copy ---------------------------------------------------
    from attribo.report.highlight import build_highlight_html

    highlight_file = tmp / "highlighted.html"
    build_highlight_html(archive_result, str(highlight_file))
    assert "Highlighted copy" in highlight_file.read_text(encoding="utf-8")
    print("[ok] highlighted copy export")

    # -- docx scan + Word review copy ---------------------------------------
    from docx import Document as DocxDocument

    from attribo.report.word_export import export_docx_with_comments

    docx_path = tmp / "chapter.docx"
    d = DocxDocument()
    d.add_paragraph(
        "Urban planning policies influence residential development patterns "
        "across metropolitan regions worldwide. " * 4
    )
    d.add_paragraph(
        "This paragraph is original writing about design studio pedagogy and "
        "contains no copied material whatsoever in this sentence."
    )
    d.save(str(docx_path))
    docx_result = run_scan(str(docx_path), archive_settings)
    review_path = tmp / "review.docx"
    added, skipped = export_docx_with_comments(docx_result, str(review_path))
    assert added >= 1, "word comment export added no comments"
    import zipfile as _zip

    with _zip.ZipFile(review_path) as zf:
        assert "word/comments.xml" in zf.namelist(), "comments part missing"
    print(f"[ok] word review copy ({added} comments, {skipped} skipped)")

    # -- batch folder scan ----------------------------------------------------
    from attribo.core.batch import run_batch

    second_doc = tmp / "second.txt"
    second_doc.write_text(
        "A short note on architectural history and modernism. " * 6, encoding="utf-8"
    )
    batch_settings = {
        "enable_web": False,
        "enable_academic": False,
        "enable_ai": True,
        "max_queries": 3,
        "min_match_words": 8,
        "exclude_quotes": False,
        "enable_archive": False,
        "fetch_timeout": 5,
        "max_fetch_per_query": 2,
        "license_key": "",
    }
    batch_out = tmp / "batch_out"
    batch_results = run_batch(
        [str(doc), str(second_doc)], batch_settings, str(batch_out)
    )
    assert len(batch_results) == 2, "batch run failed"
    assert (batch_out / "index.html").exists(), "batch index missing"
    print("[ok] batch folder scan")

    # -- PDFs ------------------------------------------------------------
    from attribo.report import pdf_builder

    report_pdf = tmp / "report.pdf"
    cert_pdf = tmp / "certificate.pdf"
    pdf_builder.build_report(result, str(report_pdf))
    pdf_builder.build_certificate(result, str(cert_pdf), licensed=False)
    assert report_pdf.stat().st_size > 2000, "report pdf too small"
    assert cert_pdf.stat().st_size > 2000, "certificate pdf too small"
    print("[ok] report + certificate PDFs")

    # -- licensing -------------------------------------------------------
    from attribo.licensing.keycheck import make_key, verify_key

    key = make_key("Test User", "2030-01-01")
    info = verify_key(key)
    assert info and info["expired"] is False
    assert verify_key("AI-BOGUS-2030-01-01-DEADBEEF00") is None
    print("[ok] licensing")

    # -- UI (offscreen) --------------------------------------------------
    from PySide6.QtWidgets import QApplication

    app = QApplication.instance() or QApplication([])
    from attribo.ui.main_window import MainWindow

    window = MainWindow(dict(settings))
    window.show()
    app.processEvents()
    window.close()
    print("[ok] UI instantiates (offscreen)")

    print("")
    print("SMOKE TEST PASSED")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
